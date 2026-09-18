from datetime import datetime

from flask import (
    Blueprint,
    flash,
    redirect,
    render_template,
    request,
    url_for,
)
from flask_login import current_user, login_required
from openpyxl import Workbook

from app.decorators import role_required
from app.extensions import db
from app.forms import AbrirTurnoForm, ArqueoForm, CompraForm, MovimientoCajaForm
from app.models.caja import (
    CategoriaMovimientoCaja,
    EstadoTurno,
    TipoMovimientoCaja,
    TurnoCaja,
)
from app.models.proveedor import Insumo, Proveedor
from app.models.usuario import Usuario
from app.services import caja_service
from app.utils.auditoria import registrar
from app.utils.excel import (
    FORMATO_MONEDA,
    ajustar_anchos,
    formatear_columna,
    marcar_encabezado,
    pesos,
    respuesta_xlsx,
)
from app.utils.moneda import parsear_centavos
from app.utils.numeros import parsear_decimal

caja_bp = Blueprint('caja', __name__, url_prefix='/caja')


def _requiere_turno_abierto():
    turno = caja_service.turno_abierto()
    if turno is None:
        flash('No hay un turno de caja abierto.', 'warning')
    return turno


@caja_bp.route('/')
@login_required
@role_required('ADMIN', 'CAJERO')
def index():
    return redirect(url_for('caja.turno'))


@caja_bp.route('/turno', methods=['GET', 'POST'])
@login_required
@role_required('ADMIN', 'CAJERO')
def turno():
    form = AbrirTurnoForm()
    actual = caja_service.turno_abierto()

    if form.validate_on_submit():
        if actual is not None:
            flash('Ya hay un turno abierto.', 'warning')
        else:
            try:
                fondo = parsear_centavos(form.fondo_inicial.data)
                caja_service.abrir_turno(fondo, current_user.id)
                db.session.commit()
                registrar('ABRIR_TURNO', 'turno_caja', detalles={'fondo_inicial': fondo})
                flash('Turno abierto.', 'success')
            except ValueError as error:
                db.session.rollback()
                flash(str(error), 'danger')
        return redirect(url_for('caja.turno'))

    if actual is None:
        return render_template('caja/turno.html', turno=None, form=form)

    return render_template(
        'caja/turno.html',
        turno=actual,
        form=form,
        kpis=caja_service.kpis_turno(actual),
        ventas_metodo=caja_service.ventas_por_metodo(actual),
        movimientos=caja_service.movimientos_ordenados(actual)[:8],
    )


@caja_bp.route('/arqueo', methods=['GET', 'POST'])
@login_required
@role_required('ADMIN', 'CAJERO')
def arqueo():
    turno_actual = caja_service.turno_abierto()
    if turno_actual is None:
        flash('No hay un turno de caja abierto.', 'warning')
        return redirect(url_for('caja.turno'))

    form = ArqueoForm()

    if form.validate_on_submit():
        try:
            contado = parsear_centavos(form.efectivo_contado.data)
            caja_service.cerrar_turno(
                turno_actual,
                contado,
                current_user.id,
                motivo_diferencia=form.motivo_diferencia.data,
                diferencia_confirmada=form.diferencia_confirmada.data,
            )
            db.session.commit()
            registrar('CERRAR_TURNO', 'turno_caja', turno_actual.id)
            flash('Turno cerrado. Se generó el cierre Z.', 'success')
            return redirect(url_for('caja.cierre_z', turno_id=turno_actual.id))
        except ValueError as error:
            db.session.rollback()
            flash(str(error), 'danger')

    esperado = caja_service.efectivo_esperado(turno_actual)
    otros = [
        v for v in caja_service.ventas_por_metodo(turno_actual)
        if not (v['metodo'] and v['metodo'].es_efectivo)
    ]
    return render_template(
        'caja/arqueo.html',
        turno=turno_actual,
        form=form,
        esperado=esperado,
        otros=otros,
    )


@caja_bp.route('/cierre-z/<int:turno_id>')
@login_required
@role_required('ADMIN', 'CAJERO')
def cierre_z(turno_id):
    turno_cerrado = db.get_or_404(TurnoCaja, turno_id)
    return render_template(
        'caja/cierre_z.html',
        turno=turno_cerrado,
        totales=caja_service.totales_cierre(turno_cerrado),
        ventas_metodo=caja_service.ventas_por_metodo(turno_cerrado),
        arqueo=turno_cerrado.arqueo,
        pedidos=caja_service.cantidad_pedidos(turno_cerrado),
    )


@caja_bp.route('/cierre-z/<int:turno_id>/exportar.xlsx')
@login_required
@role_required('ADMIN', 'CAJERO')
def cierre_z_xlsx(turno_id):
    turno_cerrado = db.get_or_404(TurnoCaja, turno_id)
    totales = caja_service.totales_cierre(turno_cerrado)
    ventas_metodo = caja_service.ventas_por_metodo(turno_cerrado)

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Cierre Z'
    hoja.append(['Cierre Z', f'Turno #{turno_cerrado.id}'])
    hoja.append(['Apertura', turno_cerrado.fecha_apertura.strftime('%d/%m/%Y %H:%M')])
    if turno_cerrado.fecha_cierre:
        hoja.append(['Cierre', turno_cerrado.fecha_cierre.strftime('%d/%m/%Y %H:%M')])
    hoja.append(['Cajero', turno_cerrado.usuario.nombre_completo])
    hoja.append([])
    hoja.append(['Concepto', 'Monto'])
    fila_conceptos = hoja.max_row

    conceptos = [
        ('Ventas', totales['ventas']),
        ('Ingresos extra', totales['ingresos']),
        ('Egresos', totales['egresos']),
        ('Total caja', totales['total_caja']),
    ]
    if turno_cerrado.arqueo:
        conceptos.append(('Diferencia arqueo', turno_cerrado.arqueo.diferencia))
    for etiqueta, monto in conceptos:
        hoja.append([etiqueta, pesos(monto)])

    hoja.append(['Pedidos', caja_service.cantidad_pedidos(turno_cerrado)])
    hoja.cell(row=hoja.max_row, column=2).number_format = '0'

    marcar_encabezado(hoja, fila=1)
    marcar_encabezado(hoja, fila=fila_conceptos)
    formatear_columna(hoja, 2, FORMATO_MONEDA, desde=fila_conceptos + 1)
    hoja.cell(row=hoja.max_row, column=2).number_format = '0'
    ajustar_anchos(hoja)

    hoja_metodos = libro.create_sheet('Ventas por método')
    hoja_metodos.append(['Método', 'Ventas', 'Monto'])
    for fila in ventas_metodo:
        hoja_metodos.append([
            fila['metodo'].nombre if fila['metodo'] else '—',
            fila['cantidad'],
            pesos(fila['monto']),
        ])
    marcar_encabezado(hoja_metodos)
    formatear_columna(hoja_metodos, 3, FORMATO_MONEDA)
    ajustar_anchos(hoja_metodos)
    hoja_metodos.freeze_panes = 'A2'

    return respuesta_xlsx(libro, f'cierre_z_{turno_id}.xlsx')


@caja_bp.route('/movimientos', methods=['GET', 'POST'])
@login_required
@role_required('ADMIN', 'CAJERO')
def movimientos():
    turno_actual = caja_service.turno_abierto()

    form = MovimientoCajaForm()
    form.proveedor_id.choices = [(0, '—')] + [
        (p.id, p.nombre) for p in Proveedor.query.order_by(Proveedor.nombre).all()
    ]

    if form.validate_on_submit():
        if turno_actual is None:
            flash('No hay un turno abierto para registrar el movimiento.', 'warning')
        else:
            try:
                monto = parsear_centavos(form.monto.data)
                caja_service.registrar_movimiento(
                    turno_actual,
                    TipoMovimientoCaja(form.tipo.data),
                    monto,
                    current_user.id,
                    categoria=(
                        CategoriaMovimientoCaja(form.categoria.data)
                        if form.categoria.data else None
                    ),
                    proveedor_id=form.proveedor_id.data or None,
                    motivo=form.motivo.data,
                )
                db.session.commit()
                registrar('MOVIMIENTO_CAJA', 'movimiento_caja', detalles={'tipo': form.tipo.data})
                flash('Movimiento registrado.', 'success')
            except ValueError as error:
                db.session.rollback()
                flash(str(error), 'danger')
        return redirect(url_for('caja.movimientos'))

    return render_template(
        'caja/movimientos.html',
        turno=turno_actual,
        form=form,
        movimientos=caja_service.movimientos_ordenados(turno_actual) if turno_actual else [],
    )


@caja_bp.route('/compras', methods=['GET', 'POST'])
@login_required
@role_required('ADMIN', 'CAJERO')
def compras():
    turno_actual = caja_service.turno_abierto()

    form = CompraForm()
    form.proveedor_id.choices = [
        (p.id, p.nombre) for p in Proveedor.query.filter_by(activo=True).order_by(Proveedor.nombre).all()
    ]

    if form.validate_on_submit():
        if turno_actual is None:
            flash('No hay un turno abierto para registrar la compra.', 'warning')
            return redirect(url_for('caja.compras'))

        try:
            lineas = _parsear_lineas_compra()
            proveedor = db.session.get(Proveedor, form.proveedor_id.data)
            caja_service.registrar_compra(
                turno_actual, proveedor, lineas, current_user.id, motivo=form.motivo.data
            )
            db.session.commit()
            registrar('COMPRA_PROVEEDOR', 'proveedor', proveedor.id, {'lineas': len(lineas)})
            flash(
                f'Compra registrada: {len(lineas)} insumo(s) y egreso en caja.',
                'success',
            )
            return redirect(url_for('caja.movimientos'))
        except (ValueError, KeyError) as error:
            db.session.rollback()
            flash(str(error), 'danger')

    return render_template(
        'caja/compras.html',
        turno=turno_actual,
        form=form,
        insumos=Insumo.query.filter_by(activo=True).order_by(Insumo.nombre).all(),
    )


def _parsear_lineas_compra():
    """Arma las líneas de compra a partir de los arrays del formulario."""
    insumo_ids = request.form.getlist('insumo_id')
    cantidades = request.form.getlist('cantidad')
    unidades = request.form.getlist('unidad')
    costos = request.form.getlist('costo')
    numeros = request.form.getlist('numero')
    vencimientos = request.form.getlist('vencimiento')

    lineas = []
    for posicion, insumo_id in enumerate(insumo_ids):
        if not insumo_id:
            continue
        insumo = db.session.get(Insumo, int(insumo_id))
        if insumo is None or not insumo.activo:
            raise ValueError('Hay un insumo inválido o inactivo en la compra.')

        cantidad = parsear_decimal(cantidades[posicion])
        if cantidad is None or cantidad <= 0:
            raise ValueError(f'Cantidad inválida para {insumo.nombre}.')

        costo = parsear_centavos(costos[posicion])

        fecha_vencimiento = datetime.strptime(vencimientos[posicion], '%Y-%m-%d').date()

        lineas.append({
            'insumo': insumo,
            'cantidad': cantidad,
            'unidad': unidades[posicion],
            'costo_unitario': costo,
            'numero': numeros[posicion].strip() or None,
            'fecha_vencimiento': fecha_vencimiento,
        })

    if not lineas:
        raise ValueError('Cargá al menos un insumo en la compra.')
    return lineas


@caja_bp.route('/historial')
@login_required
@role_required('ADMIN', 'CAJERO')
def historial():
    desde = request.args.get('desde') or ''
    hasta = request.args.get('hasta') or ''
    cajero_id = request.args.get('cajero', type=int)

    consulta = TurnoCaja.query.filter_by(estado=EstadoTurno.CERRADO)
    if desde:
        consulta = consulta.filter(
            TurnoCaja.fecha_apertura >= datetime.strptime(desde, '%Y-%m-%d')
        )
    if hasta:
        consulta = consulta.filter(
            TurnoCaja.fecha_apertura
            <= datetime.strptime(hasta, '%Y-%m-%d').replace(hour=23, minute=59)
        )
    if cajero_id:
        consulta = consulta.filter(TurnoCaja.usuario_id == cajero_id)

    turnos = consulta.order_by(TurnoCaja.fecha_apertura.desc()).all()

    cajero_ids = [
        c[0] for c in db.session.query(TurnoCaja.usuario_id).distinct().all()
    ]
    cajeros = (
        Usuario.query.filter(Usuario.id.in_(cajero_ids))
        .order_by(Usuario.apellido, Usuario.nombre)
        .all()
        if cajero_ids else []
    )

    filas = []
    for t in turnos:
        movs = t.movimientos
        filas.append({
            'turno': t,
            'ventas': sum(m.monto for m in movs if m.tipo == TipoMovimientoCaja.VENTA),
            'pedidos': caja_service.cantidad_pedidos(t),
            'diferencia': t.arqueo.diferencia if t.arqueo else 0,
        })

    return render_template(
        'caja/historial.html',
        filas=filas,
        cajeros=cajeros,
        desde=desde,
        hasta=hasta,
        cajero_id=cajero_id,
    )
