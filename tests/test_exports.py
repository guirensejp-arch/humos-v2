from io import BytesIO

from openpyxl import load_workbook

from app.models.caja import TurnoCaja


def _libro(respuesta):
    assert respuesta.status_code == 200
    return load_workbook(BytesIO(respuesta.data))


def _cerrar_turno(client):
    client.post('/caja/turno', data={'fondo_inicial': '1.000'}, follow_redirects=True)
    client.post('/caja/arqueo', data={'efectivo_contado': '1000'}, follow_redirects=True)


def test_export_inventario(client, login, datos):
    login('admin@test.com')
    libro = _libro(client.get('/inventario/exportar.xlsx'))
    assert libro.sheetnames == ['Inventario']
    assert libro.active['A1'].value == 'Insumo'


def test_export_recetas(client, login, datos):
    login('admin@test.com')
    libro = _libro(client.get('/recetas/exportar.xlsx'))
    assert libro.sheetnames == ['Recetas']
    assert libro.active['A1'].value == 'Producto'


def test_export_proveedores(client, login, datos):
    login('admin@test.com')
    libro = _libro(client.get('/proveedores/exportar.xlsx'))
    assert libro.sheetnames == ['Proveedores', 'Insumos']


def test_export_cierre_z(client, login, datos):
    login('admin@test.com')
    _cerrar_turno(client)
    turno = TurnoCaja.query.first()

    libro = _libro(client.get(f'/caja/cierre-z/{turno.id}/exportar.xlsx'))

    assert libro.sheetnames == ['Cierre Z', 'Ventas por método']


def test_export_historial(client, login, datos):
    login('admin@test.com')
    _cerrar_turno(client)

    libro = _libro(client.get('/caja/historial/exportar.xlsx?periodo=TODO'))

    assert libro.active['A1'].value == 'Turno #'
