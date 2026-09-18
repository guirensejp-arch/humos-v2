"""Sección de Analítica (solo ADMIN): KPIs, gráficos y tablas de gestión.

Independiente del Dashboard de Inicio. Es de solo lectura: arma agregados con
``analitica_service`` y exporta un resumen a Excel reutilizando ``app/utils``.
No escribe datos ni toca los módulos operativos.
"""

from flask import Blueprint, render_template, request
from flask_login import login_required
from openpyxl import Workbook

from app.decorators import admin_required
from app.services import analitica_service
from app.utils.excel import (
    FORMATO_MONEDA,
    ajustar_anchos,
    formatear_columna,
    marcar_encabezado,
    marcar_negrita,
    pesos,
    respuesta_xlsx,
)

analitica_bp = Blueprint('analitica', __name__, url_prefix='/analitica')


def _filtros():
    """Lee y valida los filtros de la URL (período, rango custom y orden)."""
    periodo = (request.args.get('periodo') or 'DIA').upper()
    if periodo not in analitica_service.PERIODOS:
        periodo = 'DIA'

    orden = (request.args.get('orden') or 'UNIDADES').upper()
    if orden not in analitica_service.ORDENES:
        orden = 'UNIDADES'

    return (
        periodo,
        request.args.get('desde') or None,
        request.args.get('hasta') or None,
        orden,
    )


def _valores_resumen(resumen):
    """Filas (etiqueta, valor, es_moneda) de los KPIs principales."""
    return [
        ('Ventas', resumen['ventas'], True),
        ('Ticket promedio', resumen['ticket_promedio'], True),
        ('Costo estimado', resumen['costo_estimado'], True),
        ('Margen bruto estimado', resumen['margen_estimado'], True),
        ('Resultado de caja', resumen['resultado_caja'], True),
        ('Pedidos', resumen['pedidos'], False),
        ('Unidades vendidas', resumen['unidades'], False),
    ]


@analitica_bp.route('/')
@login_required
@admin_required
def index():
    periodo, desde, hasta, orden = _filtros()
    datos = analitica_service.datos(periodo, desde, hasta, orden)

    return render_template(
        'analitica/index.html',
        datos=datos,
        graficos=analitica_service.graficos(datos),
        periodos=analitica_service.PERIODOS,
        ordenes=analitica_service.ORDENES,
        periodo=periodo,
        desde=desde or '',
        hasta=hasta or '',
        orden=orden,
    )


@analitica_bp.route('/exportar.xlsx')
@login_required
@admin_required
def exportar():
    periodo, desde, hasta, orden = _filtros()
    datos = analitica_service.datos(periodo, desde, hasta, orden)
    libro = Workbook()

    _hoja_resumen(libro.active, datos, periodo)
    _hoja_productos(libro.create_sheet('Productos'), datos['productos'])
    _hoja_agrupacion(libro.create_sheet('Métodos de pago'), 'Método', datos['metodos'])
    _hoja_agrupacion(libro.create_sheet('Origen'), 'Origen', datos['origenes'])
    _hoja_agrupacion(
        libro.create_sheet('Tipo de entrega'), 'Tipo de entrega', datos['entregas']
    )
    _hoja_caja(libro.create_sheet('Caja'), datos['caja'])

    return respuesta_xlsx(libro, 'analitica.xlsx')


def _hoja_resumen(hoja, datos, periodo):
    hoja.title = 'Resumen'
    hoja.append(['Analítica', datos['ventana']['etiqueta']])
    hoja.append(['Período', periodo])
    hoja.append([])
    hoja.append(['KPI', 'Valor'])
    fila_encabezado = hoja.max_row

    filas = _valores_resumen(datos['resumen'])
    for etiqueta, valor, es_moneda in filas:
        hoja.append([etiqueta, pesos(valor) if es_moneda else valor])

    marcar_encabezado(hoja, fila=fila_encabezado)
    for indice, (_, _, es_moneda) in enumerate(filas):
        if es_moneda:
            hoja.cell(row=fila_encabezado + 1 + indice, column=2).number_format = FORMATO_MONEDA
        else:
            hoja.cell(row=fila_encabezado + 1 + indice, column=2).number_format = '0'

    comparacion = datos.get('comparacion')
    if comparacion:
        hoja.append([])
        fila_comp = hoja.max_row + 1
        hoja.append(['Comparación con período anterior', 'Actual', 'Anterior', 'Variación %'])
        marcar_encabezado(hoja, fila=fila_comp)
        for clave in ('ventas', 'pedidos', 'ticket', 'unidades'):
            fila = comparacion[clave]
            es_moneda = clave in ('ventas', 'ticket')
            hoja.append([
                fila['etiqueta'],
                pesos(fila['actual']) if es_moneda else fila['actual'],
                pesos(fila['anterior']) if es_moneda else fila['anterior'],
                fila['variacion'],
            ])
            if es_moneda:
                for columna in (2, 3):
                    hoja.cell(row=hoja.max_row, column=columna).number_format = FORMATO_MONEDA

    ajustar_anchos(hoja)


def _hoja_productos(hoja, productos):
    hoja.append([
        'Producto', 'Unidades', 'Facturación', 'Precio promedio',
        'Costo estimado', 'Margen estimado', 'Participación %',
    ])
    for fila in productos:
        hoja.append([
            fila['producto'].nombre,
            fila['unidades'],
            pesos(fila['facturacion']),
            pesos(fila['precio_promedio']),
            pesos(fila['costo_estimado']),
            pesos(fila['margen_estimado']),
            fila['participacion'],
        ])
    marcar_encabezado(hoja)
    for columna in (3, 4, 5, 6):
        formatear_columna(hoja, columna, FORMATO_MONEDA)
    ajustar_anchos(hoja)
    hoja.freeze_panes = 'A2'


def _hoja_agrupacion(hoja, titulo, filas):
    hoja.append([titulo, 'Pedidos', 'Monto'])
    for fila in filas:
        hoja.append([fila['etiqueta'], fila['pedidos'], pesos(fila['monto'])])
    marcar_encabezado(hoja)
    formatear_columna(hoja, 3, FORMATO_MONEDA)
    ajustar_anchos(hoja)
    hoja.freeze_panes = 'A2'


def _hoja_caja(hoja, caja):
    hoja.append(['Concepto', 'Monto'])
    filas = [
        ('Ventas de caja', caja['ventas']),
        ('Ingresos extraordinarios', caja['ingresos']),
        ('Egresos', caja['egresos']),
        ('Resultado de caja', caja['resultado']),
        ('Compras a proveedores', caja['compras']),
        ('Gastos', caja['gastos']),
        ('Retiros del dueño', caja['retiros']),
        ('Diferencia de arqueo', caja['diferencia_arqueo']),
    ]
    for etiqueta, monto in filas:
        hoja.append([etiqueta, pesos(monto)])
    fila_resultado = 5  # encabezado + 4 conceptos (ventas, ingresos, egresos, resultado)

    if caja['egresos_detalle']:
        hoja.append([])
        hoja.append(['Egresos por categoría', 'Monto'])
        marcar_encabezado(hoja, fila=hoja.max_row)
        for fila in caja['egresos_detalle']:
            hoja.append([fila['etiqueta'], pesos(fila['monto'])])

    marcar_encabezado(hoja, fila=1)
    for fila in range(2, hoja.max_row + 1):
        celda = hoja.cell(row=fila, column=2)
        if isinstance(celda.value, (int, float)):
            celda.number_format = FORMATO_MONEDA
    marcar_negrita(hoja, fila_resultado)
    ajustar_anchos(hoja)
    hoja.freeze_panes = 'A2'
