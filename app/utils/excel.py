"""Exportación a Excel (.xlsx) con openpyxl.

Convenciones:
- Montos siempre se cargan como número (pesos = centavos / 100) con formato
  de moneda, para que Excel los sume/ordene.
- Fechas como fecha real con formato dd/mm/aaaa.
"""

from io import BytesIO

from flask import send_file
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_MONEDA = '$ #,##0.00'
FORMATO_FECHA = 'DD/MM/YYYY'

_FUENTE_ENCABEZADO = Font(bold=True, color='FFFFFF')
_FONDO_ENCABEZADO = PatternFill('solid', fgColor='263238')


def marcar_encabezado(hoja, fila=1):
    """Aplica negrita y fondo a la fila de encabezados."""
    for celda in hoja[fila]:
        if celda.value is not None:
            celda.font = _FUENTE_ENCABEZADO
            celda.fill = _FONDO_ENCABEZADO
            celda.alignment = Alignment(vertical='center')


def ajustar_anchos(hoja, minimo=8, maximo=45):
    """Ajusta el ancho de cada columna según su contenido."""
    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)
        largo = max(
            (len(str(celda.value)) for celda in columna if celda.value is not None),
            default=0,
        )
        hoja.column_dimensions[letra].width = min(max(largo + 2, minimo), maximo)


def formatear_columna(hoja, columna, formato, desde=2):
    """Aplica un formato numérico/fecha a una columna (excepto el encabezado)."""
    for fila in range(desde, hoja.max_row + 1):
        hoja.cell(row=fila, column=columna).number_format = formato


def pesos(centavos):
    """Convierte centavos a pesos (número) para celdas de Excel."""
    return round((centavos or 0) / 100, 2)


def respuesta_xlsx(libro, nombre):
    """Devuelve el libro como descarga .xlsx."""
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre,
        mimetype=(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        ),
    )
